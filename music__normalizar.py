import csv
import datetime
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

# Arquivos devem ficar na mesma pasta deste script.
PASTA = Path(__file__).resolve().parent
ARQUIVO_ENTRADA = PASTA / "music.xlsx"
ARQUIVO_SAIDA = PASTA / "music_normalizado.xlsx"

CORRECOES = {
    "Rï¿½ï¿½ma": "Rema",
    "ROSALï¿": "Rosalía",
    "Jasiel Nuï¿½ï¿½ez": "Jasiel Núñez",
    "Beï¿": "Beéle",
    "Bomba Estï¿½ï¿½reo": "Bomba Estéreo",
    "Eden Muï¿½ï": "Eden Muñoz",
    "Doï¿½ï¿½u": "Doğu Swag",
    "Beyoncï¿": "Beyoncé",
    "Zï¿½ï¿½ Neto & Crist": "Zé Neto & Cristiano",
    "Matuï¿": "Matuê",
    "Matuï¿½ï¿½": "Matuê",
    "Tiï¿½ï¿½sto": "Tiësto",
    "Tiï¿": "Tiësto",
    "Marï¿½ï¿½lia Mendonï¿½ï¿½a": "Marília Mendonça",
    "Marï¿½ï¿½lia Mendo": "Marília Mendonça",
    "Rï¿½ï": "Rema",
    "Elley Duhï¿": "Elley Duhé",
    "Michael Bublï¿": "Michael Bublé",
    "Josï¿½ï¿½ Felic": "José Feliciano",
    "Dj LK da Escï¿½ï¿½cia": "DJ LK da Escócia",
    "Mï¿½ï¿½ne": "Måneskin",
    "Xamï¿½ï¿½": "Xamã",
    "Lenny Tavï¿½ï¿½rez": "Lenny Tavárez",
    "Zï¿½ï¿½ Fe": "Zé Felipe",
    "Luï¿½ï¿½sa Sonza": "Luísa Sonza",
    "4*TOWN (From Disney and Pixarï¿½ï¿½ï¿½s Turning Red)": "4*TOWN",
    "The Marï¿": "The Marías",
    "Schï¿½ï¿½rze": "Schürze",
    "Bï¿": "BIA",
    "Cruz Cafunï¿½ï¿½": "Cruz Cafuné",
}


def corrigir_artista(nome):
    nome = nome.strip().replace("Tyler, The Creator", "Tyler_The Creator")
    for errado in sorted(CORRECOES, key=len, reverse=True):
        if nome.startswith(errado):
            return CORRECOES[errado]
    return nome


def converter_valor(valor):
    valor = valor.strip()
    if not valor or valor.lower() in {"nan", "none", "null"}:
        return None
    try:
        numero = float(valor)
        return int(numero) if numero.is_integer() else numero
    except ValueError:
        return valor


def criar_aba(wb, nome, cabecalho, linhas, nome_tabela):
    ws = wb.create_sheet(nome)
    ws.append(cabecalho)
    for linha in linhas:
        ws.append(linha)

    # Formatação simples: cabeçalho destacado, filtro e primeira linha fixa.
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")
        celula.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if ws.max_row >= 2:
        tabela = Table(displayName=nome_tabela, ref=ws.dimensions)
        tabela.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tabela)

    for coluna in ws.columns:
        letra = coluna[0].column_letter
        largura = max(len(str(c.value or "")) for c in coluna) + 2
        ws.column_dimensions[letra].width = min(max(largura, 12), 36)
    return ws


# 1. Lê a planilha original. Os registros estão armazenados como CSV na coluna A.
origem_wb = load_workbook(ARQUIVO_ENTRADA, read_only=True, data_only=True)
ws_origem = origem_wb["spotify-2023"]
linhas_csv = [str(c[0]) for c in ws_origem.iter_rows(min_col=1, max_col=1, values_only=True) if c[0]]
origem_wb.close()

cabecalho = next(csv.reader([linhas_csv[0]]))
# Remover "_%" dos nomes das colunas
cabecalho = [col.replace("_%", "") for col in cabecalho]
dados = []
for numero_linha, texto in enumerate(linhas_csv[1:], start=2):
    linha = next(csv.reader([texto]))

    # Se o título tiver vírgulas sem aspas, junta esses pedaços no primeiro campo.
    while len(linha) > len(cabecalho):
        linha[0] += "," + linha.pop(1)

    if len(linha) != len(cabecalho):
        raise ValueError(f"Linha {numero_linha}: quantidade de colunas inesperada.")
    dados.append(dict(zip(cabecalho, linha)))

# 2. Monta as três tabelas normalizadas.
colunas_musicas = ["musica_id"] + [
    coluna for coluna in cabecalho
    if coluna not in {"artist(s)_name", "artist_count"}
]
# Adicionar a coluna de data
colunas_musicas.append("release_date")

musicas = []
artistas = []
relacoes = []
ids_artistas = {}

for musica_id, registro in enumerate(dados, start=1):
    
    # Construir a nova linha
    nova_linha_musica = [musica_id]
    
    for coluna in colunas_musicas[1:-1]: # Pula 'musica_id' e 'release_date'
         nova_linha_musica.append(converter_valor(registro[coluna]))
         
    # Criar a coluna de data
    try:
        ano = int(registro.get("released_year", 0))
        mes = int(registro.get("released_month", 0))
        dia = int(registro.get("released_day", 0))
        data_lancamento = datetime.date(ano, mes, dia).strftime("%d-%m-%Y")
    except ValueError:
         data_lancamento = None

    nova_linha_musica.append(data_lancamento)

    musicas.append(nova_linha_musica)


    texto_artistas = registro["artist(s)_name"].replace(
        "Tyler, The Creator", "Tyler_The Creator"
    )
    nomes = [corrigir_artista(n) for n in texto_artistas.split(",") if n.strip()]

    for ordem, nome in enumerate(dict.fromkeys(nomes), start=1):
        if nome not in ids_artistas:
            ids_artistas[nome] = len(ids_artistas) + 1
            artistas.append([ids_artistas[nome], nome])
        relacoes.append([musica_id, ids_artistas[nome], ordem])

# 3. Salva as abas Musicas, Artistas e Musica_Artista.
saida = Workbook()
saida.remove(saida.active)

criar_aba(saida, "Musicas", colunas_musicas, musicas, "TabelaMusicas")
criar_aba(saida, "Artistas", ["artista_id", "nome_artista"], artistas, "TabelaArtistas")
criar_aba(
    saida,
    "Musica_Artista",
    ["musica_id", "artista_id", "ordem_artista"],
    relacoes,
    "TabelaMusicaArtista",
)

saida.save(ARQUIVO_SAIDA)
print(f"Arquivo criado: {ARQUIVO_SAIDA}")
print(f"Músicas: {len(musicas)} | Artistas únicos: {len(artistas)} | Relações: {len(relacoes)}")